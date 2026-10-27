# ------------------------------------------------------------------------------
# Il foglio Excel con cui si rivedono i possibili topic nuovi (vedi topic_review.py).
#
# write_review_sheet scrive un file con:
#   Istruzioni  come si compila
#   Indice      un cluster per riga, con il percorso già inserito
#   Gerarchia   i topic esistenti, per scegliere dove inserire quelli nuovi
#   C01, C02..  un foglio per cluster: in cima parole chiave, topic più simili e le
#               celle del percorso (B4:E4); sotto i progetti, dal più al meno
#               rilevante, con una colonna per escludere quelli fuori posto
#   Isolati     i progetti che non stanno in nessun cluster, ognuno con il proprio percorso
# I progetti del foglio sono già nel db con un topic provvisorio (colonna "Topic provvisorio").
# read_review_sheet legge il file compilato e restituisce {percorso: [projectID]}; controlla
# che i progetti siano ancora nel db con quel topic, così un foglio non si applica due volte
# né sopra a una decisione più recente.
# ------------------------------------------------------------------------------

import re
from collections import defaultdict

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from topic_review import (UNLABELLED, PLAIN_UNLABELLED, DEPTH, SheetError, check_name,
                          similar_topics, _cos)

ISOLATED = "Isolati"
CLUSTER_SHEET = re.compile(r"^C\d+$")
PATH_ROW = 4                  # nei fogli dei cluster, il percorso è in B4:E4
HEADER_ROW = 8                # intestazione della tabella dei progetti
FIRST_ROW = HEADER_ROW + 1
ID_COL, TOPIC_COL, OTHER_COL, ISOLATED_PATH_COL = 1, 4, 5, 5   # colonne (1-based): projectID, topic provvisorio, "Altro topic", livello 1 in Isolati
LEAF_COL = 8                  # nel foglio Gerarchia: nomi dei topic, per il menu "Altro topic"
ABSTRACT_CHARS = 600
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

INSTRUCTIONS = [
    ("Revisione dei possibili topic nuovi", True),
    ("Nell'ultimo aggiornamento di Research Intelligence alcuni progetti non sono stati inseriti in nessun cluster esistente. "
    "Alcuni di essi sono stati raggruppati in nuovi cluster (un foglio C01, C02, ... ciascuno). "
     "I progetti nel foglio Isolati non sono stati assegnati a nessun nuovo cluster. "
     "Il tuo compito è decidere se accettare i nuovi cluster, oppure se riassegnare i singoli progetti o cluster a topic esistenti. ", False),
    ("La colonna Topic provvisorio indica il topic più vicino al cluster proposto. "
     "Per gli Isolati questo valore è Unlabelled. "
     "Non modificare la colonna Topic provvisorio né i projectID.", False),
    ("", False),
    ("Come si compila", True),
    ("In ogni foglio cluster, scrivi nelle celle gialle B4:E4 il percorso del topic, dalla macro-area (livello 1) alla "
     "foglia, da 2 a 4 livelli. ", False),
    ("• Percorso vuoto: i progetti restano al topic provvisorio (diverso per ciascun progetto).", False),
    ("• Percorso già presente nel foglio Gerarchia: tutti i progetti del cluster vanno a quel topic esistente.", False),
    ("• Percorso nuovo: nasce un topic nuovo. Può stare sotto categorie esistenti o crearne di nuove. Tutti i progetti "
     "(di qualsiasi foglio) con lo stesso percorso formano un unico topic.", False),
    ("• Unlabelled nel livello 1: i progetti diventano Unlabelled.", False),
    ("", False),
    ("Dalla colonna Altro topic è possibile riassegnare un progetto di un cluster provvisorio "
     "a un cluster già esistente. Serve per assegnare eventuali progetti inseriti erroneamente nel cluster provvisorio. "
     "Si scrive o sceglie dal menu un topic esistente in cui inserire il progetto, oppure si scrive x per lasciarlo al topic "
     "provvisorio. Cancellare la riga equivale a scrivere x.", False),
    ("I progetti già presenti nella piattaforma che somigliano a un topic creato in questo foglio passano automaticamente al nuovo topic.", False),
    ("", False),
]


# ==============================================================================
# SCRITTURA
# ==============================================================================
def _put(ws, row, col, value, **style):
    """Scrive una cella come testo (un valore che inizia con "=" non diventa una formula)."""
    if isinstance(value, str):
        value = _ILLEGAL.sub("", value)
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    for key, v in style.items():
        setattr(cell, key, v)
    return cell


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _project_row(ws, row, ctx, pos, report_row, with_exclude):
    i = pos[int(report_row.projectID)]
    _put(ws, row, 1, int(report_row.projectID))
    _put(ws, row, 2, ctx.titles[i], alignment=WRAP)
    _put(ws, row, 3, ctx.objectives[i][:ABSTRACT_CHARS], alignment=WRAP)
    _put(ws, row, TOPIC_COL, str(report_row.nearest_topic), alignment=WRAP)
    if with_exclude:
        _put(ws, row, OTHER_COL, "", fill=INPUT_FILL)


def _table_header(ws, labels):
    for c, label in enumerate(labels, 1):
        _put(ws, HEADER_ROW, c, label, font=BOLD, fill=HEADER_FILL)
    ws.freeze_panes = ws.cell(row=FIRST_ROW, column=1)


def _path_cells(ws, row, first_col):
    for c in range(first_col, first_col + DEPTH):
        _put(ws, row, c, None, fill=INPUT_FILL)


def write_review_sheet(ctx, path):
    pos = {int(p): i for i, p in enumerate(ctx.project_ids)}
    report = ctx.report
    clusters = [(c, g) for c, g in report.groupby("cluster", sort=False) if c >= 0]
    isolated = report[report["cluster"] < 0]

    wb = Workbook()
    ws = wb.active
    ws.title = "Istruzioni"
    for r, (text, bold) in enumerate(INSTRUCTIONS, 1):
        _put(ws, r, 1, text, font=BOLD if bold else Font(), alignment=WRAP)
    ws.column_dimensions["A"].width = 120

    index = wb.create_sheet("Indice")
    for c, label in enumerate(["Foglio", "Progetti", "Parole chiave", "Topic esistente più simile",
                               "Percorso inserito"], 1):
        _put(index, 1, c, label, font=BOLD, fill=HEADER_FILL)
    _widths(index, [10, 10, 50, 60, 70])

    hier = wb.create_sheet("Gerarchia")
    leaves = sorted(ctx.hierarchy.leaves())
    for c, label in enumerate(["Livello 1", "Livello 2", "Livello 3", "Livello 4"], 1):
        _put(hier, 1, c, label, font=BOLD, fill=HEADER_FILL)
    for r, leaf in enumerate(leaves, 2):
        for c, node in enumerate(leaf, 1):
            _put(hier, r, c, node)
    areas = sorted({leaf[0] for leaf in leaves})
    _put(hier, 1, 6, "Macro-aree", font=BOLD, fill=HEADER_FILL)
    for r, area in enumerate(areas, 2):
        _put(hier, r, 6, area)
    leaf_names = sorted({leaf[-1] for leaf in leaves})
    _put(hier, 1, LEAF_COL, "Topic (per Altro topic)", font=BOLD, fill=HEADER_FILL)
    for r, name in enumerate(leaf_names, 2):
        _put(hier, r, LEAF_COL, name)
    _widths(hier, [34, 34, 34, 34, 4, 34, 4, 40])
    hier.freeze_panes = "A2"
    areas_range = f"Gerarchia!$F$2:$F${len(areas) + 1}"
    leaves_range = f"Gerarchia!${get_column_letter(LEAF_COL)}$2:${get_column_letter(LEAF_COL)}${len(leaf_names) + 1}"

    def add_menu(sheet, cell, source):
        """Menu a tendina che filtra mentre si digita; il testo libero resta ammesso (lo controlla la lettura)."""
        dv = DataValidation(type="list", formula1=f"={source}", allow_blank=True, showErrorMessage=False)
        sheet.add_data_validation(dv)
        dv.add(cell)

    def add_area_menu(sheet, cell):
        add_menu(sheet, cell, areas_range)

    for n, (cluster, group) in enumerate(clusters, 1):
        name = f"C{n:02d}"
        idx = [pos[int(p)] for p in group["projectID"]]
        centroid = ctx.embeddings[idx].mean(axis=0)
        order = np.argsort(_cos(ctx.embeddings[idx], centroid))[::-1]
        keywords = str(group["cluster_keywords"].iloc[0]) if group["cluster_keywords"].notna().any() else ""
        similar = similar_topics(ctx.model, ctx.hierarchy, centroid)
        similar_text = "\n".join(f"{' > '.join(p)}  ({s:.2f})" for _, p, s in similar)

        sh = wb.create_sheet(name)
        _put(sh, 1, 1, f"Cluster {n} — {len(group)} progetti", font=Font(bold=True, size=13))
        _put(sh, 2, 1, "Parole chiave", font=BOLD)
        _put(sh, 2, 2, keywords)
        _put(sh, 3, 1, "Topic esistenti più simili", font=BOLD, alignment=WRAP)
        _put(sh, 3, 2, similar_text, alignment=WRAP)
        sh.merge_cells(start_row=3, start_column=2, end_row=3, end_column=5)
        sh.row_dimensions[3].height = 15 * max(len(similar), 1)
        _put(sh, PATH_ROW, 1, "Percorso del topic", font=BOLD)
        _path_cells(sh, PATH_ROW, 2)
        add_area_menu(sh, f"B{PATH_ROW}")
        _put(sh, 5, 1, "Note", font=BOLD)
        _put(sh, 5, 2, None, fill=INPUT_FILL)
        sh.merge_cells(start_row=5, start_column=2, end_row=5, end_column=5)
        _put(sh, 6, 1, "Vuoto = resta il topic provvisorio · percorso esistente = quel topic · percorso nuovo = topic nuovo · "
                       "Unlabelled = nessun topic. Altro topic: un topic esistente per quel progetto, x = resta al topic "
                       "provvisorio.", font=Font(italic=True))
        _table_header(sh, ["projectID", "Titolo", "Objective", "Topic provvisorio", "Altro topic"])
        rows = list(group.itertuples())
        for r, k in enumerate(order, FIRST_ROW):
            _project_row(sh, r, ctx, pos, rows[k], with_exclude=True)
        add_menu(sh, f"{get_column_letter(OTHER_COL)}{FIRST_ROW}:{get_column_letter(OTHER_COL)}"
                     f"{FIRST_ROW + len(rows) - 1}", leaves_range)
        _widths(sh, [30, 60, 80, 40, 30])

        r = n + 1
        index.cell(row=r, column=1, value=f'=HYPERLINK("#\'{name}\'!A1","{name}")')
        _put(index, r, 2, len(group))
        _put(index, r, 3, keywords, alignment=WRAP)
        _put(index, r, 4, " > ".join(similar[0][1]) if similar else "", alignment=WRAP)
        levels = [f"'{name}'!{get_column_letter(c)}{PATH_ROW}" for c in range(2, 2 + DEPTH)]
        joined = f'{levels[0]}' + "".join(f'&IF({lv}="",""," > "&{lv})' for lv in levels[1:])
        index.cell(row=r, column=5, value=f'=IF({levels[0]}="","(topic provvisorio)",{joined})')
    index.freeze_panes = "A2"

    sh = wb.create_sheet(ISOLATED)
    _put(sh, 1, 1, f"Progetti isolati — {len(isolated)}", font=Font(bold=True, size=13))
    _put(sh, 2, 1, "Cosa sono: progetti troppo poco simili a tutti i topic esistenti (sotto la soglia di similarità) e "
                   "troppo diversi tra loro per formare un cluster. Per ora sono Unlabelled.")
    _put(sh, 3, 1, "Suggerimento: se un progetto non merita un topic suo, lascialo Unlabelled (cella vuota, oppure "
                   "Unlabelled nel livello 1). Per dargli un topic scrivi un percorso: lo stesso percorso nuovo su più "
                   "progetti ne forma uno solo.")
    _put(sh, 4, 1, "Vuoto = resta il topic provvisorio · percorso esistente = quel topic · percorso nuovo = topic nuovo.",
         font=Font(italic=True))
    _table_header(sh, ["projectID", "Titolo", "Objective", "Topic provvisorio",
                       "Livello 1", "Livello 2", "Livello 3", "Livello 4"])
    for r, row in enumerate(isolated.itertuples(), FIRST_ROW):
        _project_row(sh, r, ctx, pos, row, with_exclude=False)
        _path_cells(sh, r, ISOLATED_PATH_COL)
    if len(isolated):
        add_area_menu(sh, f"{get_column_letter(ISOLATED_PATH_COL)}{FIRST_ROW}:"
                          f"{get_column_letter(ISOLATED_PATH_COL)}{FIRST_ROW + len(isolated) - 1}")
    _widths(sh, [12, 60, 80, 40, 24, 24, 24, 24])

    wb.save(path)
    n_iso = len(isolated)
    print(f"Review sheet: {len(clusters)} clusters, {n_iso} isolated projects.\n"
          f"Fill in {path}, then apply it with: python update_db.py --review {path}")


# ==============================================================================
# LETTURA
# ==============================================================================
def _text(value):
    return "" if value is None else str(value).strip()


def _project_id(value):
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return int(number) if number == int(number) else None


def _read_path(levels):
    """Livelli scritti nelle celle: None se vuoti, () per Unlabelled, altrimenti la tupla.
    Solleva ValueError con il motivo se i livelli non hanno senso."""
    levels = [_text(v) for v in levels]
    filled = [v for v in levels if v]
    if not filled:
        return None
    if levels[:len(filled)] != filled:
        raise ValueError("i livelli vanno riempiti dal primo, senza buchi")
    if filled[0].casefold() == UNLABELLED.casefold():
        if len(filled) > 1:
            raise ValueError("Unlabelled non può avere altri livelli")
        return ()
    return tuple(filled)


def sheet_project_ids(path):
    """I projectID elencati nei fogli cluster e Isolati, senza altri controlli (li fa
    read_review_sheet). Serve a caricare dal db i progetti del foglio."""
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:
        raise SheetError(f"impossibile leggere {path}: {e}")
    ids = set()
    for name in wb.sheetnames:
        if CLUSTER_SHEET.match(name) or name == ISOLATED:
            ws = wb[name]
            ids |= {pid for r in range(FIRST_ROW, ws.max_row + 1)
                    if (pid := _project_id(ws.cell(r, ID_COL).value)) is not None}
    return sorted(ids)


def read_review_sheet(path, ctx):
    """Restituisce ({percorso: [projectID]}, avvisi). () è un cluster Unlabelled e
    PLAIN_UNLABELLED degli isolati Unlabelled (vedi topic_review.apply_groups); i
    progetti senza percorso o esclusi non compaiono (restano al topic provvisorio). Se c'è
    anche un solo errore solleva SheetError con l'elenco completo."""
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:
        raise SheetError(f"impossibile leggere {path}: {e}")

    errors, warnings = [], []
    known = {int(p) for p in ctx.project_ids}
    node_by_name = {n.casefold(): n for n in ctx.hierarchy.node_names()}
    leaf_paths = defaultdict(set)    # nome del topic (minuscolo) -> percorsi che lo portano
    for leaf in ctx.hierarchy.leaves():
        leaf_paths[leaf[-1].casefold()].add(leaf)

    def level(value):
        text = _text(value)
        if not text:
            return ""
        known_name = node_by_name.get(text.casefold())
        if known_name and known_name != text:
            warnings.append(f'"{text}" corretto in "{known_name}" (nome già esistente)')
        return known_name or text

    groups = defaultdict(list)       # percorso -> [(projectID, cella)]
    seen = {}
    sheets = [s for s in wb.sheetnames if CLUSTER_SHEET.match(s) or s == ISOLATED]
    if not sheets:
        errors.append("nessun foglio C01, C02, ... o Isolati")

    for name in sheets:
        ws = wb[name]
        isolated = name == ISOLATED
        cluster_path = None
        if not isolated:
            try:
                cluster_path = _read_path([level(ws.cell(PATH_ROW, c).value)
                                           for c in range(2, 2 + DEPTH)])
            except ValueError as e:
                errors.append(f"{name}!B{PATH_ROW}:E{PATH_ROW}: {e}")
        for r in range(FIRST_ROW, ws.max_row + 1):
            raw = ws.cell(r, ID_COL).value
            if _text(raw) == "":
                continue
            pid = _project_id(raw)
            if pid is None or pid not in known:
                errors.append(f"{name}!A{r}: projectID {raw!r} non presente nel db")
                continue
            if pid in seen:
                errors.append(f"{name}!A{r}: il progetto {pid} compare anche in {seen[pid]}")
                continue
            seen[pid] = f"{name}!A{r}"
            if ctx.current_names is not None:
                shown, current = _text(ws.cell(r, TOPIC_COL).value), _text(ctx.current_names.get(pid))
                if shown != current:
                    errors.append(f'{name}!A{r}: il progetto {pid} ha ora il topic "{current}" nel db, non '
                                  f'"{shown}": il foglio è già stato applicato o il progetto è stato spostato '
                                  "(cancella la riga per lasciarlo com'è)")
                    continue
            if isolated:
                try:
                    project_path = _read_path([level(ws.cell(r, ISOLATED_PATH_COL + k).value)
                                               for k in range(DEPTH)])
                except ValueError as e:
                    errors.append(f"{name}!{get_column_letter(ISOLATED_PATH_COL)}{r}: {e}")
                    continue
            else:
                other = _text(ws.cell(r, OTHER_COL).value)
                project_path = cluster_path
                if other.casefold() == "x":
                    project_path = None
                elif other:
                    found = leaf_paths.get(other.casefold(), set())
                    if len(found) == 1:
                        project_path = next(iter(found))
                    else:
                        errors.append(f'{name}!{get_column_letter(OTHER_COL)}{r}: "{other}" '
                                      + ("corrisponde a più topic" if found else "non è un topic esistente")
                                      + " (scegli dal menu, oppure x per lasciarlo al topic provvisorio)")
                        continue
            if isolated and project_path == ():
                project_path = PLAIN_UNLABELLED
            if project_path is not None:
                groups[project_path].append((pid, f"{name}!A{r}"))

    # percorsi nuovi: validati insieme, su una copia della gerarchia che li accumula
    trial = ctx.hierarchy.copy()
    free_id = max(trial.paths, default=0) + 1
    for k, new_path in enumerate(sorted((p for p in groups if p and p != PLAIN_UNLABELLED and p not in trial.leaves()), key=len)):
        where = groups[new_path][0][1]
        try:
            trial.add(free_id + k, new_path)
        except ValueError as e:
            errors.append(f"{where.split('!')[0]}: percorso {' > '.join(new_path)}: {e}")
            continue
        for node in new_path:
            if node not in node_by_name.values():
                _, similar = check_name(node, ctx.hierarchy)
                if similar:
                    warnings.append(f'"{node}" è nuovo ma simile a: {", ".join(similar)}')

    if errors:
        raise SheetError(f"{path} contains {len(errors)} error(s):\n" + "\n".join(f"  - {e}" for e in errors))
    return {p: [pid for pid, _ in ids] for p, ids in groups.items()}, list(dict.fromkeys(warnings))
